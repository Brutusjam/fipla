"""Liest das Finanzplan-Excel (Gemeinde Altdorf) und schreibt src/data.json + src/data.js.
Aufruf: python3 tools/extract_excel.py Pfad/zum/Finanzplan_BG_2027.xlsx
"""
import openpyxl, json
import sys
f=sys.argv[1] if len(sys.argv)>1 else "Finanzplan_BG_2027.xlsx"
wb=openpyxl.load_workbook(f, data_only=True)
def num(v): return round(v,4) if isinstance(v,(int,float)) else None
def rowvals(ws,r,cols): return [num(ws[f"{c}{r}"].value) for c in cols]
ER="CDEFGHIJKLMN"  # 2021..2032
out={"years":list(range(2021,2033)),"er":{},"params":{},"steuern":{},"fk":{},"invest":{"groups":[],"cat_totals":{}},"flb":{}}
ws=wb["ER Bilanz"]
for r in range(6,156):
    a=ws.cell(r,1).value; b=ws.cell(r,2).value
    vals=rowvals(ws,r,ER)
    if b is None and all(v is None for v in vals): continue
    out["er"][r]={"acc":str(a) if a is not None else "","label":(b or "").strip(),"v":vals}
ws=wb["Parameter"]
for r in list(range(5,12))+[13,14,15,16,21,24,25]+list(range(30,36))+list(range(38,50)):
    out["params"][r]={"label":(ws.cell(r,1).value or "").strip(),"v":rowvals(ws,r,ER)}
ws=wb["Steuern"]
for r in [11,12,13,15,16,17,18,21,22,23,24,29,31,32,33,34,35,36,37,40,41,42,43,44,45,46,47,48]:
    out["steuern"][r]={"label":(ws.cell(r,1).value or "").strip(),"v":rowvals(ws,r,ER)}
ws=wb["Fremdkapital"]
FK="FGHIJKLMNOPQ"
out["fk"]["bestand"]=rowvals(ws,28,FK); out["fk"]["zins_bestand"]=rowvals(ws,47,FK)
out["fk"]["loans"]=[]
for r in range(12,26):
    a=ws.cell(r,1).value
    if a: out["fk"]["loans"].append({"name":a,"zins":num(ws.cell(r,4).value),"v":rowvals(ws,r,FK)})
ws=wb["Investitionsrechnung"]
IV="DEFGHIJKLMNO"
cats=[("hochbau",11,260),("tiefbau",261,529),("mobilien",532,555),("immateriell",557,588),("informatik",589,592),("invbeitr_mit",593,615),("invbeitr_ohne",616,621),("beteiligungen",622,624)]
for cat,r0,r1 in cats:
    for r in range(r0,r1+1):
        b=ws.cell(r,2).value
        if b is None: continue
        vals=[v or 0 for v in rowvals(ws,r,IV)]
        # detail lines below until next group
        lines=[]
        rr=r+1
        while rr<=r1 and ws.cell(rr,2).value is None and ws.cell(rr,1).value is None:
            c=ws.cell(rr,3).value
            lv=[v or 0 for v in rowvals(ws,rr,IV)]
            if c and any(lv): lines.append({"name":str(c).strip(),"v":lv})
            rr+=1
        if any(vals) or lines:
            out["invest"]["groups"].append({"cat":cat,"name":str(b).strip(),"v":vals,"lines":lines})
for cat,r in zip([c[0] for c in cats],range(664,672)):
    out["invest"]["cat_totals"][cat]=[v or 0 for v in rowvals(ws,r,IV)]
out["invest"]["netto"]=[v or 0 for v in rowvals(ws,6,IV)]
ws=wb["Finanzleitbild"]
for r in range(34,43):
    out["flb"][r]={"label":(ws.cell(r,2).value or "").strip(),"v":rowvals(ws,r,ER)}
import os, sys
src=os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","src")
json.dump(out,open(os.path.join(src,"data.json"),"w"),ensure_ascii=False,indent=0)
open(os.path.join(src,"data.js"),"w").write("// Automatisch erzeugt aus dem Finanzplan-Excel durch tools/extract_excel.py\nwindow.FINANZPLAN_DATA = "+json.dumps(out,ensure_ascii=False)+";\n")
print("groups:",len(out["invest"]["groups"]),"er rows:",len(out["er"]))
# check cat totals vs group sums
for cat,_,_ in cats:
    s=[sum(g["v"][i] for g in out["invest"]["groups"] if g["cat"]==cat) for i in range(12)]
    t=out["invest"]["cat_totals"][cat]
    print(cat, [round(a-b,1) for a,b in zip(s,t)])
